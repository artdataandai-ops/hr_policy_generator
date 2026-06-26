"""
file_extract — turn an uploaded file into something a Lyzr agent can consume.

The Lyzr v3 chat endpoint takes a TEXT message, not files. So for documents we
extract plain text server-side and embed it in the message. For images there is
no text to extract, so we pass them through as base64 ``image_url`` content parts
(the same shape kyc-kyb uses) for multimodal agents.

Supported: .pdf .txt .csv .docx .pptx .xlsx .xls  → text
           .jpg .jpeg .png                          → image (base64 data URI)

extract(filename, data, content_type) -> {
  "name": str,
  "kind": "text" | "image" | "error",
  "text": str | None,         # for kind == "text"
  "data_uri": str | None,     # for kind == "image"
  "error": str | None,        # for kind == "error"
}
"""
from __future__ import annotations
import base64, csv, io, os

TEXT_EXTS = {".txt", ".csv", ".pdf", ".docx", ".pptx", ".xlsx", ".xls"}
IMAGE_EXTS = {".jpg", ".jpeg", ".png"}
IMAGE_MIME = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png"}

MAX_CHARS = 60_000  # cap extracted text per file so a giant doc can't blow the prompt


def _ext(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower()


def extract(filename: str, data: bytes, content_type: str | None = None) -> dict:
    ext = _ext(filename)
    base = {"name": filename, "kind": "error", "text": None, "data_uri": None, "error": None}
    try:
        if ext in IMAGE_EXTS:
            mime = content_type or IMAGE_MIME.get(ext, "image/jpeg")
            b64 = base64.b64encode(data).decode()
            return {**base, "kind": "image", "data_uri": f"data:{mime};base64,{b64}"}

        if ext in TEXT_EXTS:
            text = _extract_text(ext, data)
            text = (text or "").strip()
            if len(text) > MAX_CHARS:
                text = text[:MAX_CHARS] + "\n…[truncated]"
            if not text:
                return {**base, "error": "no extractable text found"}
            return {**base, "kind": "text", "text": text}

        return {**base, "error": f"unsupported file type '{ext or '?'}'"}
    except Exception as e:  # never let one bad file crash the request
        return {**base, "error": f"could not read file ({e})"}


def _extract_text(ext: str, data: bytes) -> str:
    if ext == ".txt":
        return data.decode("utf-8", errors="replace")

    if ext == ".csv":
        rows = csv.reader(io.StringIO(data.decode("utf-8", errors="replace")))
        return "\n".join(", ".join(r) for r in rows)

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)

    if ext == ".docx":
        import docx
        doc = docx.Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            parts.append(f"--- Slide {i} ---")
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    parts.append(shape.text_frame.text)
        return "\n".join(parts)

    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return _sheets_to_text(
            (ws.title, ws.iter_rows(values_only=True)) for ws in wb.worksheets
        )

    if ext == ".xls":
        import xlrd
        book = xlrd.open_workbook(file_contents=data)
        def _rows(sheet):
            for r in range(sheet.nrows):
                yield sheet.row_values(r)
        return _sheets_to_text((s.name, _rows(s)) for s in book.sheets())

    raise ValueError(f"unsupported text type {ext}")


def _sheets_to_text(sheets) -> str:
    """Render (sheet_name, rows) pairs as readable lines; rows are tuples/lists of cells."""
    out = []
    for name, rows in sheets:
        out.append(f"--- Sheet: {name} ---")
        for row in rows:
            cells = ["" if c is None else str(c) for c in row]
            if any(c.strip() for c in cells):
                out.append(", ".join(cells))
    return "\n".join(out)


def build_message(user_message: str, extracted: list[dict]) -> tuple[str, list[dict]]:
    """Combine the user's text with extracted document text, and collect image parts.

    Returns (message_text, images) where images is a list of
    {"name", "data_uri"} for any image attachments.
    """
    text_blocks, images, notes = [], [], []
    for f in extracted:
        if f["kind"] == "text":
            text_blocks.append(f"\n\n--- Attached file: {f['name']} ---\n{f['text']}")
        elif f["kind"] == "image":
            images.append({"name": f["name"], "data_uri": f["data_uri"]})
        elif f["kind"] == "error":
            notes.append(f"\n\n[Attachment '{f['name']}' could not be read: {f['error']}]")

    message = (user_message or "").strip() + "".join(text_blocks) + "".join(notes)
    return message.strip(), images
